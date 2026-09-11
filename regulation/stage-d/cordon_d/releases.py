"""Read published releases without turning row identity into real-world identity."""
from dataclasses import dataclass
import csv
import gzip
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import shapefile

from .evidence import file_digest


@dataclass(frozen=True)
class Occurrence:
    path: str
    sha256: str
    locator: str
    values: dict

    @property
    def identity(self) -> tuple[str, str]:
        return self.sha256, self.locator


def _headers(values):
    names = [str(value).strip() if value is not None else None for value in values]
    active = [name for name in names if name]
    if len(active) != len(set(active)):
        raise ValueError('Duplicate source column names require interpretation')
    return names


def workbook_occurrences(path: Path):
    """Yield every nonempty physical row, including publisher annotations.

    Formulas are retained literally. Callers must interpret them before using
    their result; this reader does not substitute cached values. Blank-header
    data is rejected rather than silently discarded. Sheets are read in full,
    regardless of an incorrect Excel dimension declaration.
    """
    digest = file_digest(path)
    # Opened as a stream: the bytes, not the file name, decide what this is.
    handle = open(path, 'rb')
    book = load_workbook(handle, read_only=True, data_only=False)
    try:
        for sheet in book:
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            names = _headers(header)
            for number, values in enumerate(rows, 2):
                if all(value is None for value in values):
                    continue
                if any(value is not None and (i >= len(names) or not names[i])
                       for i, value in enumerate(values)):
                    raise ValueError(f'Unlabelled source data: {sheet.title} row {number}')
                row = {name: values[i] if i < len(values) else None
                       for i, name in enumerate(names) if name}
                yield Occurrence(str(path), digest, f'{sheet.title}:physical-row:{number}', row)
    finally:
        book.close()
        handle.close()


def csv_occurrences(path: Path, *, encoding: str, delimiter: str):
    """CSV row numbers are logical records; embedded newlines stay in a record."""
    digest = file_digest(path)
    with path.open(encoding=encoding, newline='') as handle:
        rows = csv.reader(handle, delimiter=delimiter, strict=True)
        header = next(rows, None)
        if header is None:
            return
        names = _headers(header)
        if any(not name for name in names):
            raise ValueError('CSV needs an explicit name for every source column')
        for number, values in enumerate(rows, 2):
            if not values:
                continue
            if len(values) != len(names):
                raise ValueError(f'CSV record {number} does not match its header')
            yield Occurrence(str(path), digest, f'csv-record:{number}', dict(zip(names, values)))


def arcgis_occurrences(path: Path, *, oid_field: str, allow_repeated_oid: bool = False):
    """Read one retained native JSON chunk, preserving raw geometry and values.

    Service/layer and acquisition scope are supplied by the source manifest.
    An OID is native to that view; it is not a plant identifier. Reading a chunk
    makes no completeness claim about the service or real-world population.
    Some joined publisher views repeat OIDs. After source qualification, callers
    may explicitly retain those occurrences; their physical locators stay distinct.
    """
    digest = file_digest(path)
    data = path.read_bytes()
    if data[:2] == b'\x1f\x8b':  # gzip by its own magic, not by file name
        data = gzip.decompress(data)
    document = json.loads(data)
    if 'error' in document or 'features' not in document:
        raise ValueError('Expected an acquired ArcGIS feature response')
    seen = set()
    for number, feature in enumerate(document['features']):
        attributes = feature['attributes']
        oid = attributes.get(oid_field)
        if oid is None or (not allow_repeated_oid and (type(oid).__name__, str(oid)) in seen):
            raise ValueError('Missing or repeated native OID within a source chunk')
        seen.add((type(oid).__name__, str(oid)))
        yield Occurrence(str(path), digest, f'feature:{number}:{oid_field}:{oid}', {
            'attributes': attributes,
            'geometry': feature.get('geometry'),
            'spatialReference': document.get('spatialReference'),
        })


def shapefile_attribute_occurrences(path: Path, *, member: str, encoding: str):
    """Read native DBF facts independently of geometry availability.

    A malformed or empty source shape cannot erase its attribute record or
    subsequent records. No geometry qualification follows from this reader.
    Deleted records retain the library's omission behavior and native ordinals.
    """
    if not member.lower().endswith('.dbf'):
        raise ValueError('Expected an explicit DBF member')
    digest = file_digest(path)
    with ZipFile(path) as archive, archive.open(member) as dbf:
        with shapefile.Reader(dbf=dbf, encoding=encoding) as reader:
            for record in reader.iterRecords():
                yield Occurrence(str(path), digest, f'{member}:record:{record.oid}', {
                    'attributes': record.as_dict(),
                })


def shapefile_occurrences(path: Path, *, member: str, encoding: str):
    """Read a named ZIP member in native coordinates and source record order.

    Encoding must be established by the caller from source evidence. Projection
    text is retained, not applied; points and part offsets stay in native form.
    Deleted DBF rows follow the library's omission behavior and retain native
    record ordinals, so the output is not a claim about every physical DBF slot.
    """
    if not member.lower().endswith('.shp'):
        raise ValueError('Expected an explicit shapefile member')
    digest = file_digest(path)
    stem = member[:-4]
    with ZipFile(path) as archive:
        projection = archive.read(stem + '.prj').decode('utf-8-sig')
        with archive.open(member) as shp, archive.open(stem + '.shx') as shx, archive.open(stem + '.dbf') as dbf:
            with shapefile.Reader(shp=shp, shx=shx, dbf=dbf, encoding=encoding) as reader:
                if reader.shapeType not in {0, 1, 3, 5, 8, 11, 13, 15, 18, 21, 23, 25, 28}:
                    raise ValueError('Unsupported shapefile geometry type')
                for item in reader.iterShapeRecords():
                    yield Occurrence(str(path), digest, f'{member}:record:{item.record.oid}', {
                        'attributes': item.record.as_dict(),
                        'shape_type': item.shape.shapeType,
                        'points': list(item.shape.points),
                        'parts': list(getattr(item.shape, 'parts', ())),
                        'z': list(item.shape.z) if hasattr(item.shape, 'z') else None,
                        'm': list(item.shape.m) if hasattr(item.shape, 'm') else None,
                        'projection': projection,
                    })


def pdf_page_occurrences(path: Path):
    """Preserve each PDF page's extracted text and detected table cells.

    No fixed hash, page range, row correction or inferred column meaning. A page
    whose font cannot be decoded remains an extraction limitation, not an empty
    legal record. The original remains necessary for visual interpretation.
    """
    import pymupdf
    digest = file_digest(path)
    with pymupdf.open(path) as document:
        for index, page in enumerate(document):
            tables = [{'bbox': list(table.bbox), 'cells': table.extract()}
                      for table in page.find_tables().tables]
            yield Occurrence(str(path), digest, f'pdf-page:{index + 1}', {
                'text': page.get_text(), 'tables': tables,
                'qualification': 'Native extraction only; no semantic table interpretation or OCR completeness inferred.',
            })
